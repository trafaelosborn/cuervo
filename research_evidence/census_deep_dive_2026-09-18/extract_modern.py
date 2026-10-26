"""Read Census Bureau distribution files using their published layouts, without API keys."""
from pathlib import Path
import csv, io, json, re, zipfile
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
HERE=Path(__file__).resolve().parent
SRC=ROOT/'sources/03_census/census_deep_dive_2026-09-18'
TARGETS=['Cuervo','Santa Rosa','Guadalupe','Los Lunas','Albuquerque','Tucumcari','Vaughn','Newkirk','Montoya']

def extract2020():
    wb=openpyxl.load_workbook(SRC/'2020_dhc_geoheader.xlsx',read_only=True,data_only=True)
    fields=[str(r[1]).strip() for r in wb.active.iter_rows(values_only=True) if isinstance(r[2],(int,float)) and r[1]]
    z=zipfile.ZipFile(SRC/'nm2020.dhc.zip')
    geos=[]
    for vals in csv.reader(io.StringIO(z.read('nmgeo2020.dhc').decode('latin1')),delimiter='|'):
        assert len(vals)==len(fields),(len(vals),len(fields))
        g=dict(zip(fields,vals))
        if g['SUMLEV'] in ['040','050','060','160'] and g['GEOCOMP']=='00' and g['GEOVAR']=='00':
            if g['SUMLEV']=='040' or any(x in g['NAME'] for x in TARGETS):geos.append(g)
    w=openpyxl.load_workbook(SRC/'2020_dhc_table_matrix.xlsx',read_only=True,data_only=True)
    matrix={}
    for r in w['DHC Table Matrix'].iter_rows(values_only=True):
        if r[1] and r[2] and re.fullmatch(r'[PH][A-Z0-9]+',str(r[2]).strip()):
            matrix.setdefault(str(r[1]).strip(),[]).append({'id':str(r[2]).strip(),'label':str(r[5]),'segment':int(r[3])})
    segs={}
    for r in list(w['DHC Table_Segments'].iter_rows(values_only=True))[1:]:
        if isinstance(r[1],(int,float)):
            table=str(r[0]).strip();assert len(matrix[table])==r[2],table
            segs.setdefault(int(r[1]),[]).append((int(r[3]),table))
    selected={k for k in matrix if k in ['P1','P12','P13','P15','P16','P18','H1','H3','H4']}
    mapping={}
    for seg,ts in segs.items():
        cols=[]
        for _,table in sorted(ts):cols.extend(matrix[table])
        if any(t in selected for _,t in ts):mapping[seg]=cols
    results={g['LOGRECNO']:{'geography':g,'values':{}} for g in geos}
    for seg,cols in mapping.items():
        with z.open(f'nm{seg:05}2020.dhc') as f:
            for vals in csv.reader(io.TextIOWrapper(f,encoding='utf-8'),delimiter='|'):
                if vals[4] not in results:continue
                assert len(vals)==5+len(cols),(seg,len(vals),len(cols))
                for c,v in zip(cols,vals[5:]):
                    if c['id'] in {x['id'] for t in selected for x in matrix[t]}:
                        results[vals[4]]['values'][c['id']]={'value':v,'label':c['label'],'segment':seg}
    (HERE/'modern_2020_extract.json').write_text(json.dumps(list(results.values()),indent=2)+'\n',encoding='utf-8')
    print('2020 GEOS',[(g['NAME'],g.get('POP100'),g.get('HOUSING')) for g in geos])
    return results

def extract2010():
    # Geographic fixed widths: DPSF technical documentation, printed 2-5 and 2-6.
    z=zipfile.ZipFile(SRC/'nm2010.dp.zip');geos=[]
    for line in z.read('nmgeo2010.dp').decode('latin1').splitlines():
        g={'SUMLEV':line[8:11],'GEOCOMP':line[11:13],'LOGRECNO':line[18:25],
           'NAME':line[226:316].strip(),'STATE':line[27:29],'COUNTY':line[29:32],
           'COUSUB':line[36:41],'PLACE':line[45:50], 'POP100':line[318:327].strip(),
           'HU100':line[327:336].strip(),'AREALAND':line[198:212].strip(),'raw':line}
        if g['SUMLEV'] in ['040','050','060','160'] and g['GEOCOMP']=='00':
            if g['SUMLEV']=='040' or any(x in g['NAME'] for x in TARGETS):geos.append(g)
    # Matrix variable references occur in file order in the technical documentation.
    text=(HERE/'extracted_text/2010_dp_techdoc.txt').read_text(encoding='utf-8')
    matrix_text=text.split('TABLE (MATRIX) SECTION')[1:]
    ids=re.findall(r'dpsf\d{7}', ''.join(matrix_text))
    ids=list(dict.fromkeys(ids))
    assert len(ids)==186,len(ids)
    wanted={g['LOGRECNO']:{'geography':g,'values':{}} for g in geos}
    for row in csv.reader(io.StringIO(z.read('nm000012010.dp').decode('ascii'))):
        assert len(row)==5+len(ids),(len(row),len(ids))
        if row[4] in wanted:wanted[row[4]]['values']=dict(zip(ids,row[5:]))
    (HERE/'modern_2010_extract.json').write_text(json.dumps(list(wanted.values()),indent=2)+'\n',encoding='utf-8')
    print('2010 GEOS',[(g['NAME'],g['LOGRECNO']) for g in geos])
    return wanted

if __name__=='__main__':
    extract2020();extract2010()

